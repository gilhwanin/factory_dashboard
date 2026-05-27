# export_pan_with_mwnum.py

from UTIL.db_handler import getdb, runquery, closedb
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from datetime import datetime


# 🔹 네 PKEY 리스트 그대로 사용
PKEY_LIST = [
    26064938,26064927,26064928,26064929,26064930,26064931,26064932,26064933,
    26064934,26064935,26064936,26064937,26064939,26064940,26064941,26064942,
    26064943,26064945,26064946,26064947,26064948,26064949,26064950,26064951,
    26064952,26064953,26064954,26064955,26064956,26064957,26064958,26064959,
    26064960,26064961,26064962,26064963,26064964,26064965,26064967,26064968,
    26064969,26064970,26064971,26064972,26064973,26064974,26064975,26064976,
    26064977,26064978,26064979,26064980,26064981,26064982,26064983,26064984,
    26064985,26064986,26064987,26064988,26064989,26064990,26064991,26064992,
    26064993,26064994,26064995,26064996,26064997,26064998,26064999,26065000,
    26065002,26065003,26065004,26065005,26065006,26065007,26065008,26065009,
    26065010,26065011,26065012,26065013,26065014,26065015,26065016,26065017,
    26065018
]


def fetch_pan_by_pkeys(cur, pkey_list):
    placeholders = ",".join(["%s"] * len(pkey_list))
    query = f"""
        SELECT PKEY, BIGO, RNO, US
        FROM PAN
        WHERE PKEY IN ({placeholders})
        ORDER BY BIGO
    """
    return runquery(cur, query, tuple(pkey_list))


def process_bigo(bigo: str) -> str:
    """괄호 앞까지만 BIGO 가공"""
    if not bigo:
        return ""
    return bigo.split("(")[0].strip()


def fetch_mwnum(cur, processed_bigo: str):
    query = """
        SELECT TOP 1 MWNUM
        FROM MWNUM
        WHERE BIGO = %s
          AND RNAME = '롯데마트'
    """
    df = runquery(cur, query, (processed_bigo,))
    if df.empty:
        return None
    return df.iloc[0]["MWNUM"]


def save_to_excel(df, save_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "PAN_WITH_MWNUM"

    headers = ["PKEY", "BIGO", "RNO", "US", "MWNUM"]
    ws.append(headers)

    # 헤더 스타일
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    for _, row in df.iterrows():
        ws.append([
            row["PKEY"],
            row["BIGO"],
            row["RNO"],
            row["US"],
            row["MWNUM"]
        ])

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 32
    ws.column_dimensions["E"].width = 18

    wb.save(save_path)


def main():
    conn, cur = getdb("GYUN_N")
    if not conn:
        print("❌ DB 연결 실패")
        return

    try:
        print("✅ DB 연결 성공")

        pan_df = fetch_pan_by_pkeys(cur, PKEY_LIST)

        if pan_df.empty:
            print("PAN 조회 결과 없음")
            return

        # MWNUM 컬럼 추가
        mwnum_list = []

        for _, row in pan_df.iterrows():
            processed = process_bigo(row["BIGO"])
            mwnum = fetch_mwnum(cur, processed)
            mwnum_list.append(mwnum)

        pan_df["MWNUM"] = mwnum_list

        today = datetime.now().strftime("%Y%m%d")
        save_path = f"C:\\P\\lotte_pan_with_mwnum_{today}.xlsx"

        save_to_excel(pan_df, save_path)

        print("\n===== 엑셀 저장 완료 =====")
        print("저장경로:", save_path)
        print("총 행 수:", len(pan_df))

    finally:
        closedb(conn)


if __name__ == "__main__":
    main()
